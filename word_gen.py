"""
word_gen.py — Report generation orchestrator.
Change this file only if the ORDER or SET of steps in generate_report needs
to change. All implementation lives in the other word_gen_*.py modules.

Module map:
  word_gen_core.py         low-level XML/run helpers
  word_gen_replacements.py build_replacements() — all {{placeholder}} values
  word_gen_contacts.py     contact, notification, signature, occupancy tables
  word_gen_matrix.py       integrations matrix, gen-served expansion, diagram
  word_gen_sections.py     test procedure sections, absent-section removal
  word_gen_appendix.py     Appendix B (results) and Appendix C (blank form)
"""

import os
import zipfile
import shutil

from docx import Document
from lxml import etree

from word_gen_core import (
    replace_all,
    replace_placeholder_paragraph_with_paragraphs,
    remove_system_section,
    style_arencon_runs,
)
from word_gen_replacements import build_replacements
from word_gen_contacts import (
    populate_contacts_table,
    populate_notifications_table,
    populate_signatures_table,
    populate_occupancies_table,
    populate_ist_notes_table,
)
from word_gen_matrix import (
    populate_matrix_table,
    expand_gen_served_system,
    insert_diagram_image,
)
from word_gen_sections import (
    populate_test_procedures,
    _remove_absent_tp_blocks,
    expand_gen_served_tp,
    populate_elevator_action_table,
)
from word_gen_appendix import (
    _save_appc_template_rows,
    populate_appendix_b_table,
    populate_appendix_c_table,
)
from constants import SYSTEMS

# Re-export pick_date so existing callers (ui_data.py etc.) keep working
from word_gen_core import pick_date  # noqa: F401


def generate_report(data, template_path, output_path):
    doc = Document(template_path)

    # Matrix table must be populated BEFORE replace_all so placeholders are clean
    populate_matrix_table(doc, data)

    # Expand gen_served_system/list BEFORE replace_all wipes the placeholders
    expand_gen_served_system(doc, data)

    # Save Appendix C template rows BEFORE replace_all wipes their placeholders
    appc_template_rows = _save_appc_template_rows(doc)

    replace_all(doc, build_replacements(data))

    # Building description
    building_desc = data.get("building_description", "")
    if building_desc:
        paras = [p.strip() for p in building_desc.split("\n\n") if p.strip()]
        replace_placeholder_paragraph_with_paragraphs(doc, "building_description", paras)

    # Personnel Safety (Section 5)
    ps = data.get("personnel_safety", {})
    for key in ("safety_protocols", "special_hazards", "team_communications", "occupant_notification"):
        text = ps.get(key, "")
        if text:
            paras = [p.strip() for p in text.split("\n") if p.strip()]
            replace_placeholder_paragraph_with_paragraphs(doc, key, paras)

    # Systems
    systems_data = data.get("systems", {})
    replacements = build_replacements(data)
    for sys_info in SYSTEMS:
        key = sys_info["key"]
        sys = systems_data.get(key, {})
        if sys.get("present", False):
            for text_key, ph_key in [("description", "desc_ph"), ("integrations", "int_ph")]:
                text = sys.get(text_key, "")
                if text:
                    # Substitute any inline placeholders (e.g. {{facp_room}}) before inserting
                    for ph_key2, val in replacements.items():
                        text = text.replace(f"{{{{{ph_key2}}}}}", str(val) if val else "")
                    paras = [p.strip() for p in text.split("\n") if p.strip()]
                    replace_placeholder_paragraph_with_paragraphs(doc, sys_info[ph_key], paras)
        else:
            remove_system_section(doc, [sys_info["desc_ph"], sys_info["int_ph"]])

    # Contractors
    contractors = data.get("contractors", [])
    if contractors:
        populate_contacts_table(doc, contractors)
        populate_notifications_table(doc, contractors)
        populate_signatures_table(doc, contractors)

    # Occupancies
    occupancies = data.get("occupancies", [])
    if occupancies:
        populate_occupancies_table(doc, occupancies)

    # IST Notes (Appendix B)
    ist_notes = data.get("ist_notes", [])
    if ist_notes:
        populate_ist_notes_table(doc, ist_notes)

    # Expand generator served TP bullets BEFORE test procedures run
    expand_gen_served_tp(doc, data)

    # Test procedures — clone heading+table per integration row
    populate_test_procedures(doc, data, replacements)

    # Remove leftover TP template blocks for absent systems
    _remove_absent_tp_blocks(doc, data)

    # Appendix B — Integrated Systems Testing Results table
    populate_appendix_b_table(doc, data)

    # Appendix C — Blank integrated testing form (same structure, no results/notes)
    populate_appendix_c_table(doc, data, appc_template_rows)

    # Interconnection diagram image
    diag_png = data.get("diagram_png")
    if diag_png and os.path.exists(diag_png):
        insert_diagram_image(doc, diag_png)

    # Elevator action table (Action | Operation Description)
    populate_elevator_action_table(doc, data)

    # Style all "Arencon Inc." occurrences in body/tables (headers excluded)
    style_arencon_runs(doc)

    doc.save(output_path)

    # Inject updateFields into settings.xml so Word refreshes the TOC on open
    _set_update_fields(output_path)


def _set_update_fields(output_path):
    """Re-open the saved docx and inject <w:updateFields w:val='true'/> into
    word/settings.xml so Word prompts to refresh the TOC when the file is opened."""
    settings_path = "word/settings.xml"
    tmp_path = output_path + ".tmp"
    WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    W = f"{{{WNS}}}"
    try:
        with zipfile.ZipFile(output_path, "r") as zin, \
                zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data_bytes = zin.read(item.filename)
                if item.filename == settings_path:
                    root = etree.fromstring(data_bytes)
                    existing = root.find(f"{W}updateFields")
                    if existing is None:
                        uf = etree.Element(f"{W}updateFields")
                        uf.set(f"{W}val", "true")
                        root.insert(0, uf)
                    else:
                        existing.set(f"{W}val", "true")
                    data_bytes = etree.tostring(root, xml_declaration=True,
                                                encoding="UTF-8", standalone=True)
                zout.writestr(item, data_bytes)
        shutil.move(tmp_path, output_path)
    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)